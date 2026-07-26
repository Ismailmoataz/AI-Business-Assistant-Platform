"""Tenant-safe document storage, extraction, search, and lifecycle workflows."""

from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from uuid import UUID, uuid4

from docx import Document as DocxDocument
from fastapi import UploadFile
from minio.error import S3Error
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import and_, func, or_, select
from sqlalchemy.orm import Session, selectinload

from app.config.settings import settings
from app.database.models.business import Activity, Customer, Tag
from app.database.models.documents import Document, DocumentFavorite, DocumentRecent, DocumentVersion, Folder
from app.database.models.enums import ActivityAction, DocumentVisibility, UserRole
from app.database.models.user import User
from app.services.business_service import BusinessConflictError, BusinessNotFoundError, BusinessPermissionError, BusinessValidationError
from app.services.infrastructure import get_minio_client

ALLOWED_MIME_TYPES = {
    "application/pdf", "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "text/plain", "text/markdown", "text/csv", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "image/jpeg", "image/png", "image/gif", "image/webp",
}


class DocumentService:
    """Coordinates authorization, object storage, extraction, and document history."""

    def __init__(self, session: Session, storage=None) -> None:
        """Bind the service to one transaction and optionally injected storage client."""
        self.session, self.storage = session, storage or get_minio_client()

    def _activity(self, actor: User, action: ActivityAction, entity: object, summary: str) -> None:
        """Record a Phase 3 mutation in the existing tenant activity history."""
        self.session.add(Activity(company_id=actor.company_id, actor_id=actor.id, action=action, entity_type=entity.__class__.__name__.lower(), entity_id=entity.id, summary=summary))

    def _commit(self, item):
        """Commit and refresh one entity with conflict translation."""
        try: self.session.commit()
        except Exception as exc:
            self.session.rollback()
            if "unique" in str(exc).lower(): raise BusinessConflictError("A record with those values already exists") from exc
            raise
        self.session.refresh(item); return item

    def _visible(self, actor: User, item) -> bool:
        """Evaluate private, department, and company access without cross-tenant leakage."""
        if item.company_id != actor.company_id: return False
        if actor.role == UserRole.ADMIN: return True
        if item.visibility == DocumentVisibility.COMPANY: return True
        if item.visibility == DocumentVisibility.DEPARTMENT: return actor.department_id is not None and item.department_id == actor.department_id
        return item.owner_id == actor.id

    def _visibility_filter(self, actor: User):
        """Build the database predicate matching the document visibility policy."""
        if actor.role == UserRole.ADMIN:
            return Document.company_id == actor.company_id
        return and_(
            Document.company_id == actor.company_id,
            or_(
                Document.visibility == DocumentVisibility.COMPANY,
                Document.owner_id == actor.id,
                and_(
                    Document.visibility == DocumentVisibility.DEPARTMENT,
                    Document.department_id == actor.department_id,
                    actor.department_id is not None,
                ),
            ),
        )

    def _manage(self, actor: User, item) -> bool:
        """Return whether the actor can change a folder or document."""
        return actor.role == UserRole.ADMIN or item.owner_id == actor.id or (actor.role == UserRole.MANAGER and item.department_id == actor.department_id)

    def _document(self, actor: User, document_id: UUID, manage=False) -> Document:
        """Return an active document only if the requested access is authorized."""
        item = self.session.scalar(select(Document).options(selectinload(Document.tags), selectinload(Document.versions)).where(Document.id == document_id, Document.company_id == actor.company_id, Document.deleted_at.is_(None)))
        if item is None or not self._visible(actor, item) or (manage and not self._manage(actor, item)): raise BusinessNotFoundError("Document was not found")
        return item

    def _folder(self, actor: User, folder_id: UUID, manage=False) -> Folder:
        """Return an active folder only if the requested access is authorized."""
        item = self.session.scalar(select(Folder).where(Folder.id == folder_id, Folder.company_id == actor.company_id, Folder.deleted_at.is_(None)))
        if item is None or not self._visible(actor, item) or (manage and not self._manage(actor, item)): raise BusinessNotFoundError("Folder was not found")
        return item

    def _customer(self, actor: User, customer_id: UUID) -> Customer:
        """Resolve an active tenant customer for document metadata links."""
        if actor.role == UserRole.EMPLOYEE:
            raise BusinessPermissionError("Employees cannot link customer records")
        customer = self.session.scalar(
            select(Customer).where(Customer.id == customer_id, Customer.company_id == actor.company_id, Customer.deleted_at.is_(None))
        )
        if customer is None:
            raise BusinessValidationError("Customer was not found")
        return customer

    def _extract(self, data: bytes, mime: str) -> str | None:
        """Extract text locally from supported non-image formats; images remain metadata-only."""
        if mime.startswith("image/"): return None
        if mime in {"text/plain", "text/markdown", "text/csv"}: return data.decode("utf-8", errors="replace")
        if mime == "application/pdf": return "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(data)).pages)
        if mime.endswith("wordprocessingml.document"):
            return "\n".join(paragraph.text for paragraph in DocxDocument(BytesIO(data)).paragraphs)
        if mime.endswith("spreadsheetml.sheet"):
            workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
            return "\n".join("\t".join("" if value is None else str(value) for row in sheet.iter_rows(values_only=True) for value in row) for sheet in workbook.worksheets)
        return None

    def _put(self, key: str, data: bytes, mime: str) -> None:
        """Ensure the private bucket exists and upload one immutable object."""
        bucket = settings.minio_documents_bucket
        if not self.storage.bucket_exists(bucket):
            try:
                self.storage.make_bucket(bucket)
            except S3Error as exc:
                if exc.code not in {"BucketAlreadyExists", "BucketAlreadyOwnedByYou"}:
                    raise
        self.storage.put_object(bucket, key, BytesIO(data), len(data), content_type=mime)

    def create_folder(self, actor: User, name: str, parent_id: UUID | None, department_id: UUID | None, visibility: DocumentVisibility) -> Folder:
        """Create a scoped folder after validating parent and department boundaries."""
        if visibility != DocumentVisibility.PRIVATE and actor.role == UserRole.EMPLOYEE: raise BusinessPermissionError("Employees may create private folders only")
        if visibility == DocumentVisibility.DEPARTMENT and department_id != actor.department_id and actor.role != UserRole.ADMIN: raise BusinessPermissionError("Department is outside your scope")
        if parent_id: self._folder(actor, parent_id, manage=True)
        folder = Folder(company_id=actor.company_id, owner_id=actor.id, name=name, parent_id=parent_id, department_id=department_id, visibility=visibility)
        self.session.add(folder); self.session.flush(); self._activity(actor, ActivityAction.CREATED, folder, "Created folder"); return self._commit(folder)

    async def upload(self, actor: User, upload: UploadFile, title: str, description: str | None, folder_id: UUID | None, department_id: UUID | None, visibility: DocumentVisibility, tag_ids: list[UUID], customer_id: UUID | None = None) -> Document:
        """Validate, store, extract, and version a new document without exposing storage credentials."""
        mime = upload.content_type or "application/octet-stream"
        if mime not in ALLOWED_MIME_TYPES: raise BusinessValidationError("Unsupported file type")
        if visibility != DocumentVisibility.PRIVATE and actor.role == UserRole.EMPLOYEE: raise BusinessPermissionError("Employees may upload private documents only")
        if visibility == DocumentVisibility.DEPARTMENT and department_id != actor.department_id and actor.role != UserRole.ADMIN: raise BusinessPermissionError("Department is outside your scope")
        if customer_id: self._customer(actor, customer_id)
        if folder_id: self._folder(actor, folder_id, manage=True)
        data = await upload.read()
        if not data or len(data) > settings.document_max_upload_bytes: raise BusinessValidationError("File is empty or exceeds the upload limit")
        tags = list(self.session.scalars(select(Tag).where(Tag.company_id == actor.company_id, Tag.id.in_(tag_ids), Tag.deleted_at.is_(None)))) if tag_ids else []
        if len(tags) != len(set(tag_ids)): raise BusinessValidationError("One or more tags were not found")
        document_id, version = uuid4(), 1; key = f"{actor.company_id}/{document_id}/{version}/{Path(upload.filename or 'upload').name}"
        self._put(key, data, mime); text = self._extract(data, mime)
        doc = Document(id=document_id, company_id=actor.company_id, owner_id=actor.id, folder_id=folder_id, department_id=department_id, customer_id=customer_id, title=title.strip(), filename=Path(upload.filename or "upload").name, mime_type=mime, size_bytes=len(data), description=description, visibility=visibility, object_key=key, extracted_text=text, tags=tags)
        self.session.add(doc); self.session.flush(); self.session.add(DocumentVersion(document_id=doc.id, uploaded_by_user_id=actor.id, version=1, filename=doc.filename, mime_type=mime, size_bytes=len(data), object_key=key, extracted_text=text)); self._activity(actor, ActivityAction.CREATED, doc, "Uploaded document"); self._commit(doc)
        return doc

    async def replace(self, actor: User, document_id: UUID, upload: UploadFile) -> Document:
        """Store a replacement object as the next immutable document version."""
        doc = self._document(actor, document_id, manage=True); mime = upload.content_type or "application/octet-stream"
        if mime not in ALLOWED_MIME_TYPES: raise BusinessValidationError("Unsupported file type")
        data = await upload.read()
        if not data or len(data) > settings.document_max_upload_bytes: raise BusinessValidationError("File is empty or exceeds the upload limit")
        version = doc.version + 1; filename = Path(upload.filename or doc.filename).name; key = f"{actor.company_id}/{doc.id}/{version}/{filename}"
        self._put(key, data, mime); text = self._extract(data, mime)
        doc.version, doc.filename, doc.mime_type, doc.size_bytes, doc.object_key, doc.extracted_text = version, filename, mime, len(data), key, text
        self.session.add(DocumentVersion(document_id=doc.id, uploaded_by_user_id=actor.id, version=version, filename=filename, mime_type=mime, size_bytes=len(data), object_key=key, extracted_text=text)); self._activity(actor, ActivityAction.UPDATED, doc, "Replaced document"); return self._commit(doc)

    def update(self, actor: User, document_id: UUID, changes: dict) -> Document:
        """Apply authorized metadata, move, tag, or permission changes to a document."""
        doc = self._document(actor, document_id, manage=True)
        if "folder_id" in changes and changes["folder_id"]: self._folder(actor, changes["folder_id"], manage=True)
        if "visibility" in changes and changes["visibility"] != DocumentVisibility.PRIVATE and actor.role == UserRole.EMPLOYEE: raise BusinessPermissionError("Employees may not change shared permissions")
        tag_ids = changes.pop("tag_ids", None)
        if changes.get("customer_id"): self._customer(actor, changes["customer_id"])
        for key, value in changes.items(): setattr(doc, key, value)
        if tag_ids is not None:
            tags = list(self.session.scalars(select(Tag).where(Tag.company_id == actor.company_id, Tag.id.in_(tag_ids), Tag.deleted_at.is_(None)))) if tag_ids else []
            if len(tags) != len(set(tag_ids)): raise BusinessValidationError("One or more tags were not found")
            doc.tags = tags
        self._activity(actor, ActivityAction.UPDATED, doc, "Updated document metadata"); return self._commit(doc)

    def list_documents(self, actor: User, offset: int, limit: int, query: str | None, folder_id: UUID | None, owner_id: UUID | None = None, customer_id: UUID | None = None) -> tuple[list[Document], int]:
        """Search visible documents across metadata, text, tags, and folder names."""
        statement = select(Document).options(selectinload(Document.tags)).where(self._visibility_filter(actor), Document.deleted_at.is_(None))
        if query:
            pattern = f"%{query.strip()}%"
            statement = statement.outerjoin(Document.tags).outerjoin(Document.folder).where(or_(Document.title.ilike(pattern), Document.filename.ilike(pattern), Document.description.ilike(pattern), Document.extracted_text.ilike(pattern), Tag.name.ilike(pattern), Folder.name.ilike(pattern)))
        if folder_id: statement = statement.where(Document.folder_id == folder_id)
        if owner_id: statement = statement.where(Document.owner_id == owner_id)
        if customer_id: statement = statement.where(Document.customer_id == customer_id)
        count_statement = select(func.count()).select_from(statement.order_by(None).with_only_columns(Document.id).distinct().subquery())
        total = self.session.scalar(count_statement) or 0
        items = self.session.scalars(statement.order_by(Document.updated_at.desc()).offset(offset).limit(limit)).unique().all()
        return list(items), total

    def record_download(self, actor: User, document_id: UUID, version: int | None = None) -> tuple[Document, DocumentVersion]:
        """Authorize a version download and record a recent view plus activity event."""
        doc = self._document(actor, document_id); selected = next((v for v in doc.versions if v.version == (version or doc.version)), None)
        if selected is None: raise BusinessNotFoundError("Document version was not found")
        recent = self.session.scalar(select(DocumentRecent).where(DocumentRecent.user_id == actor.id, DocumentRecent.document_id == doc.id))
        if recent: recent.updated_at = datetime.now(UTC)
        else: self.session.add(DocumentRecent(user_id=actor.id, document_id=doc.id))
        self._activity(actor, ActivityAction.UPDATED, doc, "Downloaded document"); self.session.commit(); return doc, selected

    def soft_delete(self, actor: User, document_id: UUID, restore: bool = False) -> Document:
        """Soft-delete or restore a document under management policy."""
        doc = self._document(actor, document_id, manage=True) if not restore else self.session.scalar(select(Document).where(Document.id == document_id, Document.company_id == actor.company_id))
        if doc is None or not self._manage(actor, doc): raise BusinessNotFoundError("Document was not found")
        doc.deleted_at = None if restore else datetime.now(UTC); self._activity(actor, ActivityAction.UPDATED if restore else ActivityAction.DELETED, doc, "Restored document" if restore else "Deleted document"); return self._commit(doc)

    def list_folders(self, actor: User) -> list[Folder]:
        """List visible active folders in the current tenant."""
        visibility = Folder.company_id == actor.company_id if actor.role == UserRole.ADMIN else and_(
            Folder.company_id == actor.company_id,
            or_(Folder.visibility == DocumentVisibility.COMPANY, Folder.owner_id == actor.id, and_(Folder.visibility == DocumentVisibility.DEPARTMENT, Folder.department_id == actor.department_id, actor.department_id is not None)),
        )
        return list(self.session.scalars(select(Folder).where(visibility, Folder.deleted_at.is_(None)).order_by(Folder.name)))

    def update_folder(self, actor: User, folder_id: UUID, changes: dict) -> Folder:
        """Rename, move, or adjust visibility for an authorized folder."""
        folder = self._folder(actor, folder_id, manage=True)
        if changes.get("parent_id"):
            parent = self._folder(actor, changes["parent_id"], manage=True)
            if parent.id == folder.id: raise BusinessValidationError("A folder cannot be its own parent")
        if changes.get("visibility") != DocumentVisibility.PRIVATE and changes.get("visibility") and actor.role == UserRole.EMPLOYEE: raise BusinessPermissionError("Employees may not share folders")
        for key, value in changes.items(): setattr(folder, key, value)
        self._activity(actor, ActivityAction.UPDATED, folder, "Updated folder"); return self._commit(folder)

    def favorite(self, actor: User, document_id: UUID, enabled: bool) -> None:
        """Add or remove a user-specific document favorite."""
        doc = self._document(actor, document_id); favorite = self.session.scalar(select(DocumentFavorite).where(DocumentFavorite.user_id == actor.id, DocumentFavorite.document_id == doc.id))
        if enabled and favorite is None: self.session.add(DocumentFavorite(user_id=actor.id, document_id=doc.id))
        if not enabled and favorite: self.session.delete(favorite)
        self.session.commit()

    def favorites(self, actor: User) -> list[Document]:
        """Return the caller's still-accessible favorite documents."""
        return list(self.session.scalars(select(Document).join(DocumentFavorite).options(selectinload(Document.tags)).where(DocumentFavorite.user_id == actor.id, self._visibility_filter(actor), Document.deleted_at.is_(None))).all())

    def recents(self, actor: User) -> list[Document]:
        """Return the caller's recently downloaded/accessed documents."""
        return list(self.session.scalars(select(Document).join(DocumentRecent).options(selectinload(Document.tags)).where(DocumentRecent.user_id == actor.id, self._visibility_filter(actor), Document.deleted_at.is_(None)).order_by(DocumentRecent.updated_at.desc())).all())
